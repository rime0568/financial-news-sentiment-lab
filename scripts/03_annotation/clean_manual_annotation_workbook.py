"""Extract manual labels from the annotated workbook and restore canonical text.

The annotated workbook contains duplicate label columns because Excel inserted
the labels near the front while retaining the original blank columns. This
script reads the first label columns, joins canonical text from the V2 unified
dataset, and writes a new UTF-8-BOM CSV without touching the workbook.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


MOJIBAKE_RE = re.compile(r"(?:�|锟|鈥|тА|тШ|Ã|Â|бн)")


def find_workbook(directory: Path) -> Path:
    files = [p for p in directory.glob("*manually*.xlsx") if not p.name.startswith("~$")]
    if not files:
        raise FileNotFoundError("No manually annotated workbook found")
    return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)[0]


def read_annotations(workbook_path: Path) -> pd.DataFrame:
    ws = load_workbook(workbook_path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    # The workbook has duplicate label headers. The first six columns are the
    # user-entered annotation columns; the later duplicate columns are blank.
    records = []
    for row in rows[1:]:
        records.append(
            {
                "annotation_id": row[0],
                "source": row[1],
                "ticker": row[2],
                "human_sentiment": row[3],
                "annotation_confidence": row[4],
                "annotation_note": row[5],
                "company_name": row[6],
                "title_from_workbook": row[7],
                "clean_text_from_workbook": row[8],
                "publication_datetime": row[9],
                "url": row[10],
                "accession_number": row[11],
            }
        )
    return pd.DataFrame(records)


def build(workbook_path: Path, canonical_path: Path, output_path: Path, report_path: Path) -> None:
    annotations = read_annotations(workbook_path)
    annotations["human_sentiment"] = annotations["human_sentiment"].fillna("").astype(str).str.strip().str.lower()
    annotations["annotation_confidence"] = annotations["annotation_confidence"].fillna("").astype(str).str.strip().str.lower()
    canonical = pd.read_csv(canonical_path, low_memory=False)
    canonical["ticker"] = canonical["ticker"].astype(str).str.upper().str.strip()
    annotations["ticker"] = annotations["ticker"].astype(str).str.upper().str.strip()

    key = ["source", "ticker", "title_from_workbook"]
    canonical_key = canonical[["source", "ticker", "title", "company_name", "clean_text", "publication_datetime", "url", "accession_number"]].copy()
    canonical_key = canonical_key.rename(
        columns={
            "title": "title_from_workbook",
            "company_name": "company_name_canonical",
            "clean_text": "clean_text_canonical",
            "publication_datetime": "publication_datetime_canonical",
            "url": "url_canonical",
            "accession_number": "accession_number_canonical",
        }
    )
    merged = annotations.merge(canonical_key, on=key, how="left", suffixes=("", "_canonical"))
    merged["title"] = merged["title_from_workbook"]
    merged["clean_text"] = merged["clean_text_canonical"].fillna(merged["clean_text_from_workbook"])
    merged["company_name"] = merged["company_name_canonical"].fillna(merged["company_name"])
    merged["publication_datetime"] = merged["publication_datetime_canonical"].fillna(merged["publication_datetime"])
    merged["url"] = merged["url_canonical"].fillna(merged["url"])
    merged["accession_number"] = merged["accession_number_canonical"].fillna(merged["accession_number"])
    def quality_flags(row: pd.Series) -> str:
        flags = []
        if row["source"] == "sec_edgar_8k":
            flags.append("sec_full_filing_not_headline")
        if bool(MOJIBAKE_RE.search(str(row["clean_text"]))):
            flags.append("possible_mojibake")
        if pd.isna(row["clean_text_canonical"]):
            flags.append("canonical_text_not_matched")
        return ";".join(flags) if flags else "ok"

    merged["text_quality_flag"] = merged.apply(quality_flags, axis=1)

    output_columns = [
        "annotation_id", "source", "ticker", "human_sentiment", "annotation_confidence", "annotation_note",
        "company_name", "title", "clean_text", "publication_datetime", "url", "accession_number", "text_quality_flag",
    ]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    merged[output_columns].to_csv(output_path, index=False, encoding="utf-8-sig")

    flags = merged["text_quality_flag"].fillna("")
    report_rows = [
        {"metric": "rows", "value": len(merged)},
        {"metric": "manual_labels_filled", "value": int(merged["human_sentiment"].fillna("").astype(str).str.strip().ne("").sum())},
        {"metric": "positive_labels", "value": int(merged["human_sentiment"].eq("positive").sum())},
        {"metric": "neutral_labels", "value": int(merged["human_sentiment"].eq("neutral").sum())},
        {"metric": "negative_labels", "value": int(merged["human_sentiment"].eq("negative").sum())},
        {"metric": "sec_full_filing_rows", "value": int(flags.str.contains("sec_full_filing_not_headline", regex=False).sum())},
        {"metric": "possible_mojibake_rows", "value": int(flags.str.contains("possible_mojibake", regex=False).sum())},
        {"metric": "canonical_text_not_matched", "value": int(flags.str.contains("canonical_text_not_matched", regex=False).sum())},
    ]
    pd.DataFrame(report_rows).to_csv(report_path, index=False, encoding="utf-8-sig")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=None)
    parser.add_argument("--annotation-dir", type=Path, default=Path("data/interim/annotation"))
    parser.add_argument("--canonical", type=Path, default=Path("data/processed_v2/unified_text_dataset.csv"))
    parser.add_argument("--output", type=Path, default=Path("data/interim/annotation/manual_annotations_cleaned.csv"))
    parser.add_argument("--report", type=Path, default=Path("outputs/model_evaluation/manual_annotation_quality.csv"))
    args = parser.parse_args()
    workbook = args.workbook or find_workbook(args.annotation_dir)
    build(workbook, args.canonical, args.output, args.report)
